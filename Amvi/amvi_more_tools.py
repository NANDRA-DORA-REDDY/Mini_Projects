"""more_tools.py - Amvi v7 tools: code-to-file, vision, study helper, git, custom apps.
Imported by jarvis.py; it registers its tools into the same TOOLS registry as tools.py."""
import base64, io, json, os, re, shutil, subprocess, sys, time
from datetime import datetime
from pathlib import Path

import requests

import amvi_tools as tools
from amvi_tools import tool, WORKSPACE, OLLAMA_URL, CODER_MODEL, FALLBACK_MODEL, APPS, _mem, model_for

CODE_DIR = WORKSPACE / "code"
PROJECT_DIR = WORKSPACE / "projects"
STUDY_DIR = WORKSPACE / "study"
for _d in (CODE_DIR, PROJECT_DIR, STUDY_DIR):
    _d.mkdir(parents=True, exist_ok=True)

STUDY_MODEL = FALLBACK_MODEL        # llama3.2 is good enough for notes/MCQs
VISION_MODEL = "moondream"          # ollama pull moondream  (small; or use llava)


INPUT = input   # the GUI replaces this with a pop-up


def _generate(prompt, model, images=None, json_mode=False, timeout=300):
    payload = {"model": model, "prompt": prompt, "stream": False, "options": {"num_ctx": 8192}}
    if images: payload["images"] = images
    if json_mode: payload["format"] = "json"
    r = requests.post(OLLAMA_URL, json=payload, timeout=timeout)
    if r.status_code == 404:
        raise RuntimeError(f"Model '{model}' not found. Run: ollama pull {model}")
    r.raise_for_status()
    return r.json().get("response", "").strip()


def _code_llm(prompt):
    try:
        return _generate(prompt, model_for("coder", CODER_MODEL))
    except RuntimeError:
        return _generate(prompt, FALLBACK_MODEL)


def _find(name, *bases):
    p = Path(os.path.expandvars(str(name).strip().strip('"\'')))
    if p.is_absolute():
        return p
    for b in (bases or (WORKSPACE,)):
        if (b / p).exists():
            return b / p
    return (bases[0] if bases else WORKSPACE) / p


# =============== 1. CODE TO FILE, RUN, PROJECTS ===============
LANGS = {"python": ".py", "py": ".py", "c": ".c", "c++": ".cpp", "cpp": ".cpp",
         "javascript": ".js", "js": ".js", "java": ".java", "html": ".html"}


@tool("Write a complete program and SAVE it to a file (python, c, c++, javascript, java, html)",
      '{"description": str, "language": str, "filename": str}')
def write_code(description, language="python", filename=""):
    ext = LANGS.get(language.lower(), ".py")
    raw = _code_llm(f"Write a complete, working {language} program: {description}\n"
                    "Return ONLY the code in one block, with short comments.")
    m = re.search(r"```[\w+#-]*\n(.*?)```", raw, re.S)
    code = (m.group(1) if m else raw).strip() + "\n"
    if not filename:
        filename = re.sub(r"\W+", "_", description.lower()).strip("_")[:30] or "program"
    if not filename.lower().endswith(ext):
        filename += ext
    p = CODE_DIR / filename
    if p.exists():
        p = p.with_name(f"{p.stem}_{int(time.time())}{p.suffix}")
    p.write_text(code, encoding="utf-8")
    if shutil.which("code"):
        subprocess.Popen(f'code "{p}"', shell=True)
    return f"Saved {p}\nSay: run {p.name}\n\n{code[:600]}"


@tool("Run a saved program (.py .js .c .cpp) and show its output", '{"path": str}', dangerous=True)
def run_code(path):
    p = _find(path, CODE_DIR, WORKSPACE)
    if not p.exists():
        return f"File not found: {p}"
    ext = p.suffix.lower()
    if ext == ".py":
        cmd = [sys.executable, str(p)]
    elif ext == ".js":
        if not shutil.which("node"): return "Node.js is not installed."
        cmd = ["node", str(p)]
    elif ext in (".c", ".cpp"):
        cc = shutil.which("gcc" if ext == ".c" else "g++")
        if not cc: return "No C/C++ compiler found (install MinGW-w64 and add it to PATH)."
        exe = p.with_suffix(".exe")
        c = subprocess.run([cc, str(p), "-o", str(exe)], capture_output=True, text=True)
        if c.returncode: return "Compile error:\n" + c.stderr[:1500]
        cmd = [str(exe)]
    else:
        return f"I can't run {ext} files."
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=20, cwd=p.parent, input="")
    except subprocess.TimeoutExpired:
        return "Stopped: the program ran longer than 20 seconds."
    out = (r.stdout + ("\n[errors]\n" + r.stderr if r.stderr else "")).strip()
    return out[:2000] or "(no output)"


@tool("Review a code file: find bugs, explain them and show a fix", '{"path": str, "question": str}')
def review_code(path, question=""):
    p = _find(path, CODE_DIR, WORKSPACE)
    if not p.exists():
        return f"File not found: {p}"
    code = p.read_text(encoding="utf-8", errors="ignore")[:6000]
    return _code_llm(f"{question or 'Find any bugs or problems in this code, explain them, and show the corrected code.'}"
                     f"\n\nFile {p.name}:\n{code}")


@tool("Create a project structure: kind is python, web, c or cpp", '{"name": str, "kind": str}')
def create_project(name, kind="python"):
    safe = re.sub(r"[^\w\- ]", "", name).strip().replace(" ", "_") or "project"
    root = PROJECT_DIR / safe
    kind = kind.lower().replace("c++", "cpp").replace("html", "web")
    files = {
        "python": {"main.py": 'def main():\n    print("Hello from ' + safe + '")\n\n\nif __name__ == "__main__":\n    main()\n',
                   "requirements.txt": "", "README.md": f"# {safe}\n", ".gitignore": ".venv/\n__pycache__/\n"},
        "web": {"index.html": '<!DOCTYPE html>\n<html>\n<head>\n  <meta charset="UTF-8">\n  <title>' + safe +
                '</title>\n  <link rel="stylesheet" href="style.css">\n</head>\n<body>\n  <h1>' + safe +
                '</h1>\n  <script src="script.js"></script>\n</body>\n</html>\n',
                "style.css": "body { font-family: sans-serif; margin: 2rem; }\n", "script.js": "console.log('ready');\n"},
        "c": {"main.c": '#include <stdio.h>\n\nint main(void) {\n    printf("Hello\\n");\n    return 0;\n}\n'},
        "cpp": {"main.cpp": '#include <iostream>\n\nint main() {\n    std::cout << "Hello" << std::endl;\n    return 0;\n}\n'},
    }.get(kind)
    if files is None:
        return "Kinds: python, web, c, cpp."
    root.mkdir(parents=True, exist_ok=True)
    for fn, content in files.items():
        if not (root / fn).exists():
            (root / fn).write_text(content, encoding="utf-8")
    tools.remember(f"{safe.replace('_', ' ')} project", str(root))
    return f"Created {kind} project at {root}\nSay: open my {safe.replace('_', ' ')} project"


@tool("Install a Python package with pip", '{"package": str}', dangerous=True)
def pip_install(package):
    if not re.fullmatch(r"[A-Za-z0-9_.\-\[\]=<>]+", package):
        return "That package name looks invalid."
    r = subprocess.run([sys.executable, "-m", "pip", "install", package],
                       capture_output=True, text=True, timeout=300)
    return (r.stdout + r.stderr)[-500:]


# =============== 2. VISION ===============
def _vision(image, question):
    image = image.convert("RGB")
    image.thumbnail((1280, 1280))
    buf = io.BytesIO(); image.save(buf, "PNG")
    try:
        return _generate(question or "Describe what you see in detail.", model_for("vision", VISION_MODEL),
                         images=[base64.b64encode(buf.getvalue()).decode()])
    except RuntimeError as e:
        return str(e)


@tool("Describe or answer a question about what is on the screen using a vision model",
      '{"question": str, "delay": number}')
def describe_screen(question="", delay=4):
    from PIL import ImageGrab
    delay = max(0, min(int(delay), 15))
    if delay:
        print(f"Amvi: Switch to the window you want me to look at... capturing in {delay}s")
        time.sleep(delay)
    return _vision(ImageGrab.grab(), question)


@tool("Describe or answer a question about an image file (charts, diagrams, photos)",
      '{"path": str, "question": str}')
def describe_image(path, question=""):
    from PIL import Image
    p = _find(path, WORKSPACE, STUDY_DIR)
    if not p.exists():
        return f"File not found: {p}"
    return _vision(Image.open(p), question)


# =============== 3. STUDY ASSISTANT ===============
DOC_EXTS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv", ".py", ".c", ".cpp", ".js"}
TASKS = {
    "summary": "Summarize the material clearly in simple language with the key points.",
    "notes": "Make well-organized study notes with headings and bullet points.",
    "mcq": "Create 10 multiple-choice questions with 4 options each (A-D), then an answer key.",
    "2mark": "Write 5 likely 2-mark exam questions, each with a short 2-3 line answer.",
    "4mark": "Write 4 likely 4-mark exam questions, each with an answer of about 8-10 lines.",
    "8mark": "Write 2 likely 8-mark exam questions with detailed structured answers "
             "(definition, explanation, example, conclusion).",
}


def _doc_text(p):
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if ext == ".docx":
        import docx
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(p)).slides, 1):
            out.append(f"--- Slide {i} ---")
            out += [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        return "\n".join(out)
    return p.read_text(encoding="utf-8", errors="ignore")


@tool("Study helper. task: summary, notes, mcq, 2mark, 4mark or 8mark. source: a file path "
      "(pdf, docx, pptx, txt) or a topic", '{"task": str, "source": str}')
def study(task, source):
    t = re.sub(r"[\s_\-]", "", task.lower())
    key = ("mcq" if "mcq" in t or "multiple" in t else "2mark" if "2" in t else "4mark" if "4" in t
           else "8mark" if "8" in t else "notes" if "note" in t else "summary")
    p = _find(source, STUDY_DIR, WORKSPACE)
    note = ""
    if p.suffix.lower() in DOC_EXTS:
        if not p.exists():
            return f"File not found: {p}"
        try:
            material = _doc_text(p)
        except ImportError as e:
            return f"Missing package ({e.name}). Run: pip install -r requirements.txt"
        if not material.strip():
            return "I couldn't extract any text (scanned PDF? try read image / OCR)."
        if len(material) > 12000:
            material, note = material[:12000], "\n\n(Note: only the first part of a long file was used.)"
        label = p.stem
    else:
        material, label = f"Topic: {source}\n(Use your own knowledge.)", re.sub(r"\W+", "_", source)[:30]
    out = _generate(f"{TASKS[key]}\n\nMaterial:\n{material}", model_for("study", STUDY_MODEL))
    save = STUDY_DIR / f"{key}_{label}.txt"
    save.write_text(out, encoding="utf-8")
    return f"{out}{note}\n\n[Saved to {save}]"


@tool("Interactive multiple-choice quiz on a topic", '{"topic": str, "count": number}')
def quiz_me(topic, count=5):
    try:
        raw = _generate(f'Create {int(count)} multiple-choice questions about "{topic}". Return JSON: '
                        '{"questions":[{"q":"...","options":["...","...","...","..."],"answer":0}]} '
                        "where answer is the index (0-3) of the correct option.",
                        model_for("study", STUDY_MODEL), json_mode=True)
        qs = json.loads(raw)["questions"]
    except Exception as e:
        return f"I couldn't build the quiz ({e}). Try again."
    score = asked = 0
    for q in qs:
        try:
            opts, ans = q["options"], int(q["answer"])
            assert 2 <= len(opts) <= 4 and 0 <= ans < len(opts)
        except Exception:
            continue
        asked += 1
        print(f"\nQ{asked}. {q['q']}")
        for i, o in enumerate(opts):
            print(f"  {'ABCD'[i]}) {o}")
        a = INPUT("Your answer (A-D, or q to stop): ").strip().upper()
        if a == "Q":
            asked -= 1
            break
        if a and a[0] in "ABCD"[:len(opts)] and "ABCD".index(a[0]) == ans:
            score += 1; print("  Correct!")
        else:
            print(f"  Not quite. Answer: {'ABCD'[ans]}) {opts[ans]}")
    log_study(topic, 0)
    return f"Quiz finished: {score}/{asked}." if asked else "Quiz stopped."


@tool("Log study time on a topic", '{"topic": str, "minutes": number}')
def log_study(topic, minutes=0):
    f = STUDY_DIR / "log.json"
    data = json.loads(f.read_text()) if f.exists() else []
    data.append({"date": datetime.now().strftime("%Y-%m-%d"), "topic": topic, "minutes": float(minutes)})
    f.write_text(json.dumps(data, indent=2))
    return f"Logged {minutes} min on {topic}."


@tool("Show study progress by topic")
def study_report():
    f = STUDY_DIR / "log.json"
    if not f.exists():
        return "No study sessions logged yet."
    totals = {}
    for e in json.loads(f.read_text()):
        t = totals.setdefault(e["topic"], [0, 0.0]); t[0] += 1; t[1] += e["minutes"]
    return "\n".join(f"{k}: {v[0]} sessions, {v[1]:.0f} min" for k, v in totals.items())


# =============== 4. GIT ===============
def _git(args, repo="", timeout=60):
    if not shutil.which("git"):
        return "Git is not installed (get it from git-scm.com)."
    v = _mem().get(repo.lower(), repo) if repo else _mem().get("amvi project", "")
    p = Path(v) if v else WORKSPACE
    if not p.is_absolute():
        p = PROJECT_DIR / p
    if not p.exists():
        return f"Folder not found: {p}"
    r = subprocess.run(["git", *args], cwd=p, capture_output=True, text=True, timeout=timeout)
    return f"[{p}]\n" + ((r.stdout + r.stderr).strip()[:1500] or "(done)")


@tool("Git status of a project", '{"repo": str}')
def git_status(repo=""): return _git(["status", "-sb"], repo)


@tool("Git recent commits", '{"repo": str}')
def git_log(repo=""): return _git(["log", "--oneline", "-8"], repo)


@tool("Commit all changes with a message", '{"message": str, "repo": str}', dangerous=True)
def git_commit(message="Update via Amvi", repo=""):
    _git(["add", "-A"], repo)
    return _git(["commit", "-m", message], repo)


@tool("Push commits to GitHub", '{"repo": str}', dangerous=True)
def git_push(repo=""): return _git(["push"], repo, timeout=120)


@tool("Pull latest changes", '{"repo": str}', dangerous=True)
def git_pull(repo=""): return _git(["pull"], repo, timeout=120)


@tool("Switch to a git branch (creates it if missing)", '{"name": str, "repo": str}')
def git_branch(name, repo=""):
    if not re.fullmatch(r"[\w./\-]+", name):
        return "Invalid branch name."
    out = _git(["switch", name], repo)
    return out if "fatal" not in out and "invalid" not in out.lower() else _git(["switch", "-c", name], repo)


@tool("Clone a GitHub repository (https URL)", '{"url": str, "folder": str}', dangerous=True)
def git_clone(url, folder=""):
    if not url.startswith("https://"):
        return "Only https:// URLs are allowed."
    dest = PROJECT_DIR / (folder or url.rstrip("/").split("/")[-1].removesuffix(".git"))
    r = subprocess.run(["git", "clone", url, str(dest)], capture_output=True, text=True, timeout=300)
    return (r.stdout + r.stderr).strip()[-800:] or f"Cloned to {dest}"


# =============== 5. YOUR OWN APPS AND GAMES ===============
APPS_FILE = WORKSPACE / "apps.json"


def _load_apps():
    if APPS_FILE.exists():
        for name, path in json.loads(APPS_FILE.read_text()).items():
            APPS[name] = (f'"{path}"', Path(path).name)


_load_apps()


@tool("Teach Amvi an app or game: a name and the full path to its .exe", '{"name": str, "path": str}',
      dangerous=True)
def add_app(name, path):
    p = Path(path.strip().strip('"\''))
    if not p.exists():
        return f"Path not found: {p}"
    data = json.loads(APPS_FILE.read_text()) if APPS_FILE.exists() else {}
    data[name.lower()] = str(p)
    APPS_FILE.write_text(json.dumps(data, indent=2))
    APPS[name.lower()] = (f'"{p}"', p.name)
    return f"Learned '{name}'. Say: open {name}"


@tool("List the apps Amvi can open")
def list_apps(): return ", ".join(sorted(APPS))


# =============== 6. PLAY ON YOUTUBE ===============
@tool("Find a video or song on YouTube and play it right away", '{"query": str}')
def play_youtube(query):
    if not query.strip():
        return "What should I play?"
    ids = []
    try:
        r = requests.get("https://www.youtube.com/results",
                         params={"search_query": query, "sp": "EgIQAQ=="},   # videos only
                         headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                                                "(KHTML, like Gecko) Chrome/120.0 Safari/537.36",
                                  "Accept-Language": "en-US,en;q=0.9"}, timeout=10)
        ids = re.findall(r'"videoId":"([\w-]{11})"', r.text)
    except Exception:
        pass
    if ids:
        tools.open_chrome(f"https://www.youtube.com/watch?v={ids[0]}&autoplay=1")
        return f"Playing the top YouTube result for '{query}'. (Say 'pause' or 'next' to control it.)"
    return "Couldn't pick a video automatically. " + tools.search_site("youtube", query)


# =============== 7. HELP ===============
HELP_TEXT = r"""Amvi commands (examples). Exact patterns are instant; other phrasing goes through the AI.

Web:       open youtube | search for python tutorials | search github for ai assistant
Media:     play lofi music on youtube | pause | resume | next song | volume up | mute
Apps:      open vs code | open calculator | close notepad | restart chrome | list apps
           add app NAME at C:\path\app.exe
Windows:   lock the computer | shutdown | restart the computer | sleep | take a screenshot
Files:     create a folder called X | find my python files | disk space | open D:\path
System:    how much ram am I using | what is my gpu usage | battery status
Memory:    remember that my X is at PATH | open my X | add a note: text
Timers:    start a 30 minute study timer | remind me to study at 7 pm | show reminders
           (reminders: "in 30 minutes", "at 19:30", "at 7 pm", "tomorrow at 9 am")
Study:     summarize file.pdf | make mcqs on TOPIC | make 4 mark answers from file.pdf
           quiz me on TOPIC | study report | add deadline: TASK on 2026-12-01 | show my deadlines
Coding:    write a python program to ... | run file.py | review file.py
           create a python project called X | paste (for long multi-line questions)
Git:       git status | commit: message | push this project | git pull | switch branch NAME
Screen:    read my screen | explain the error on my screen | describe my screen | read image FILE
Setups:    start my coding setup (menu) | start my study setup | full coding setup
Menus:     open app | web search | coding setup   (buttons in the GUI)
Multi-step: join commands with 'and' / 'then': open chrome and search for python
Voice:     voice on | voice off | listen | voice list | voice speed 170
           voice neerja (neural, Indian English) | voice shruti (Telugu) | voice swara (Hindi) | voice local
New (v11): hands-free (GUI checkbox) then say: Hey Amvi, <command>  |  stop talking
           minimize window | maximize window | switch window | show desktop | close window
           explain clipboard | summarize clipboard | fix clipboard | analyze data.csv
           summarize https://site.com/page
New (v12): good morning (daily briefing) | personality funny | start with windows | voice wake add <word>
Other:     help | exit"""


@tool("Show the list of commands Amvi really supports")
def show_help():
    return HELP_TEXT + f"\n\n({len(tools.TOOLS)} tools loaded.)"


# =============== 8. REMINDERS (saved to disk, survive restarts) ===============
import threading
from datetime import timedelta

REMINDERS_FILE = WORKSPACE / "reminders.json"
_rem_lock = threading.Lock()


def _load_rem():
    return json.loads(REMINDERS_FILE.read_text()) if REMINDERS_FILE.exists() else []


def _parse_when(s):
    s = s.lower().strip()
    now = datetime.now()
    m = re.fullmatch(r"in\s+(\d+(?:\.\d+)?)\s*(seconds?|secs?|minutes?|mins?|hours?|hrs?|days?)", s)
    if m:
        n, u = float(m.group(1)), m.group(2)[0]
        return now + {"s": timedelta(seconds=n), "m": timedelta(minutes=n),
                      "h": timedelta(hours=n), "d": timedelta(days=n)}[u]
    tomorrow = s.startswith("tomorrow")
    if tomorrow:
        s = re.sub(r"^tomorrow\s*(?:at)?\s*", "", s) or "9:00"
    m = re.fullmatch(r"(?:at\s+)?(\d{1,2})(?::(\d{2}))?\s*(a\.?m\.?|p\.?m\.?)?", s)
    if not m:
        return None
    h, mi, ap = int(m.group(1)), int(m.group(2) or 0), (m.group(3) or "").replace(".", "")
    if ap == "pm" and h < 12: h += 12
    if ap == "am" and h == 12: h = 0
    if h > 23 or mi > 59:
        return None
    due = datetime.combine(now.date(), datetime.min.time()).replace(hour=h, minute=mi)
    if tomorrow:
        due += timedelta(days=1)
    elif due <= now:
        due += timedelta(days=1)          # that time has passed today: use tomorrow
    return due


@tool("Set a reminder that is saved to disk. when: 'in 30 minutes', 'at 7 pm', 'at 19:30', 'tomorrow at 9 am'",
      '{"text": str, "when": str}')
def remind_me(text, when):
    due = _parse_when(when)
    if not due:
        return "I couldn't understand the time. Try: in 30 minutes | at 7 pm | at 19:30 | tomorrow at 9 am"
    with _rem_lock:
        data = _load_rem()
        data.append({"text": text, "due": due.isoformat(timespec="seconds"), "done": False})
        REMINDERS_FILE.write_text(json.dumps(data, indent=2))
    return f"Reminder set for {due.strftime('%a %d %b, %I:%M %p')}: {text}"


@tool("Show pending reminders")
def show_reminders():
    pend = sorted((r for r in _load_rem() if not r["done"]), key=lambda r: r["due"])
    return "\n".join(f"{datetime.fromisoformat(r['due']).strftime('%a %d %b %I:%M %p')}  {r['text']}"
                     for r in pend) or "No pending reminders."


@tool("Clear all reminders")
def clear_reminders():
    with _rem_lock:
        REMINDERS_FILE.write_text("[]")
    return "All reminders cleared."


def _reminder_loop():
    while True:
        try:
            with _rem_lock:
                data, changed = _load_rem(), False
                for r in data:
                    if not r["done"] and datetime.fromisoformat(r["due"]) <= datetime.now():
                        r["done"], changed = True, True
                        print(f"\n[Amvi] \u23f0 Reminder: {r['text']}\nYOU: ", end="", flush=True)
                        try:
                            import winsound; winsound.Beep(1000, 400)
                        except Exception:
                            pass
                if changed:
                    REMINDERS_FILE.write_text(json.dumps(data, indent=2))
        except Exception:
            pass
        time.sleep(15)


threading.Thread(target=_reminder_loop, daemon=True).start()


# =============== 9. LONG-TERM MEMORY (facts + chat history that survive restarts) ===============
FACTS_FILE = WORKSPACE / "facts.json"
HISTORY_FILE = WORKSPACE / "chat_history.json"


def _facts():
    return json.loads(FACTS_FILE.read_text()) if FACTS_FILE.exists() else []


def facts_text():
    f = _facts()[-15:]
    return "\n".join(f"- {x}" for x in f) if f else "(none)"


@tool("Remember a fact about the user or their life (exams, preferences, plans)", '{"fact": str}')
def remember_fact(fact):
    facts = _facts()
    if fact.strip() and fact.strip() not in facts:
        facts.append(fact.strip())
        FACTS_FILE.write_text(json.dumps(facts, indent=2, ensure_ascii=False), encoding="utf-8")
    return f"I'll remember that: {fact.strip()}"


@tool("Forget remembered facts that contain a keyword", '{"keyword": str}')
def forget_fact(keyword):
    keyword = re.sub(r"^(?:the|my|about)\s+", "", keyword.strip(), flags=re.I)
    facts = _facts()
    keep = [x for x in facts if keyword.lower() not in x.lower()]
    FACTS_FILE.write_text(json.dumps(keep, indent=2, ensure_ascii=False), encoding="utf-8")
    return f"Forgot {len(facts) - len(keep)} fact(s)."


@tool("Show what Amvi remembers about the user")
def show_facts():
    return facts_text()


def load_history(n=20):
    try:
        return [tuple(x) for x in json.loads(HISTORY_FILE.read_text(encoding="utf-8"))][-n:]
    except Exception:
        return []


def save_history(history):
    try:
        HISTORY_FILE.write_text(json.dumps(history[-50:], ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass


# =============== 10. MODELS PER ROLE ===============
EMBED_MODEL = "nomic-embed-text"                    # ollama pull nomic-embed-text
ROLE_DEFAULTS = {"router": FALLBACK_MODEL, "coder": CODER_MODEL, "study": STUDY_MODEL,
                 "vision": VISION_MODEL, "agent": CODER_MODEL, "embed": EMBED_MODEL}


def _installed():
    r = requests.get(OLLAMA_URL.replace("/api/generate", "/api/tags"), timeout=5)
    return [m["name"] for m in r.json().get("models", [])]


@tool("Show installed AI models and which model each role uses")
def show_models():
    try:
        names = _installed()
    except Exception as e:
        return f"Can't reach Ollama: {e}"
    roles = "\n".join(f"  {r:7} -> {model_for(r, d)}" for r, d in ROLE_DEFAULTS.items())
    return "Installed: " + ", ".join(names) + "\n\nRoles:\n" + roles + \
           "\n\nChange one with: use coder model qwen2.5-coder:7b"


@tool("Choose the model for a role: router, coder, study, vision, agent or embed", '{"role": str, "name": str}')
def set_model(role, name):
    role = role.lower()
    if role not in ROLE_DEFAULTS:
        return "Roles: " + ", ".join(ROLE_DEFAULTS)
    try:
        names = _installed()
    except Exception as e:
        return f"Can't reach Ollama: {e}"
    if name not in names and f"{name}:latest" not in names:
        return f"'{name}' is not installed. Run: ollama pull {name}\nInstalled: {', '.join(names)}"
    data = json.loads(SETTINGS_FILE.read_text()) if SETTINGS_FILE.exists() else {}
    data.setdefault("models", {})[role] = name
    SETTINGS_FILE.write_text(json.dumps(data, indent=2))
    return f"The {role} model is now {name}."


# =============== 11. DOCUMENT SEARCH (local RAG) ===============
KB_DIR = WORKSPACE / "kb"
KB_DIR.mkdir(parents=True, exist_ok=True)
KB_FILE = KB_DIR / "index.json"


def _embed(text):
    r = requests.post(OLLAMA_URL.replace("/api/generate", "/api/embeddings"),
                      json={"model": model_for("embed", EMBED_MODEL), "prompt": text}, timeout=60)
    if r.status_code == 404:
        raise RuntimeError(f"Embedding model not found. Run: ollama pull {EMBED_MODEL}")
    r.raise_for_status()
    return r.json()["embedding"]


def _chunks(text, size=900, overlap=150):
    text = re.sub(r"\s+", " ", text).strip()
    return [text[i:i + size] for i in range(0, len(text), size - overlap)][:400]


@tool("Index the study folder (or a folder/file) so Amvi can answer questions from your documents",
      '{"folder": str}')
def index_docs(folder=""):
    root = _find(folder, STUDY_DIR, WORKSPACE) if folder else STUDY_DIR
    files = [p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in DOC_EXTS] if root.is_dir() \
        else [root]
    idx = json.loads(KB_FILE.read_text(encoding="utf-8")) if KB_FILE.exists() else {"files": {}, "chunks": []}
    new = 0
    for p in files:
        if idx["files"].get(str(p)) == p.stat().st_mtime:
            continue
        print(f"Indexing {p.name}...")
        try:
            text = _doc_text(p)
            pieces = [(i, c, _embed(c)) for i, c in enumerate(_chunks(text))]
        except Exception as e:
            print(f"Skipped {p.name}: {e}")
            continue
        idx["chunks"] = [c for c in idx["chunks"] if c["file"] != str(p)]
        idx["chunks"] += [{"file": str(p), "name": p.name, "i": i, "text": c, "vec": v} for i, c, v in pieces]
        idx["files"][str(p)] = p.stat().st_mtime
        new += 1
    KB_FILE.write_text(json.dumps(idx), encoding="utf-8")
    return f"Indexed {new} new/changed file(s). {len(idx['files'])} files, {len(idx['chunks'])} chunks in total."


@tool("Answer a question using ONLY the user's indexed documents (notes, PDFs, slides)", '{"question": str}')
def ask_docs(question):
    import numpy as np
    if not KB_FILE.exists():
        return "Nothing is indexed yet. Say: index my documents"
    chunks = json.loads(KB_FILE.read_text(encoding="utf-8"))["chunks"]
    if not chunks:
        return "Nothing is indexed yet. Put files in D:\\AmviWorkspace\\study and say: index my documents"
    q = np.array(_embed(question))
    mat = np.array([c["vec"] for c in chunks])
    sims = mat @ q / (np.linalg.norm(mat, axis=1) * np.linalg.norm(q) + 1e-9)
    top = sims.argsort()[::-1][:4]
    context = "\n\n".join(f"[{chunks[i]['name']}]\n{chunks[i]['text']}" for i in top)
    ans = _generate("Answer the question using ONLY the context below. If the answer is not in the context, "
                    f"say you could not find it.\n\nContext:\n{context}\n\nQuestion: {question}",
                    model_for("study", STUDY_MODEL))
    return f"{ans}\n\n[Sources: {', '.join(sorted({chunks[i]['name'] for i in top}))}]"


# =============== 12. AUTONOMOUS AGENT ===============
import inspect
import logging

CONFIRM = input                         # the GUI replaces this with a pop-up
AGENT_STOP = threading.Event()          # the GUI's Stop button sets this
MAX_STEPS = 8
AGENT_BLOCKED = {"shutdown_computer", "restart_computer", "sleep_computer", "delete_file", "pip_install",
                 "git_push", "git_clone", "add_app", "clear_reminders", "run_agent", "quiz_me", "show_help",
                 "close_window", "startup_launch"}


def _agent_llm(prompt, model):
    try:
        return _generate(prompt, model, json_mode=True, timeout=180)
    except RuntimeError:
        return _generate(prompt, FALLBACK_MODEL, json_mode=True, timeout=180)


@tool("Autonomous agent: give a goal and Amvi plans and runs several tools, step by step", '{"goal": str}')
def run_agent(goal):
    AGENT_STOP.clear()
    model = model_for("agent", CODER_MODEL)
    catalog = "\n".join(f'- {n}: {t["desc"]} args={t["args"] or "{}"}'
                        for n, t in tools.TOOLS.items() if n not in AGENT_BLOCKED)
    system = ("You are the planner inside Amvi, a Windows assistant. Achieve the user's goal by using the "
              "tools below ONE step at a time.\nReply with ONE JSON object:\n"
              '{"thought": "short reasoning", "action": "tool_name", "args": {...}}  to use a tool, or\n'
              '{"final": "message for the user"}  when the goal is done or cannot be done.\n'
              "Rules: use only the listed tools and exact argument names; never repeat a step; be efficient; "
              "for questions you can answer yourself, reply with final.\n\n"
              f"Known facts about the user:\n{facts_text()}\n\nTools:\n{catalog}\n")
    obs, seen, bad = [], set(), 0
    for n in range(1, MAX_STEPS + 1):
        if AGENT_STOP.is_set():
            return "Agent stopped.\n" + "\n".join(obs)
        steps = "\n".join(obs) or "(none yet)"
        try:
            d = json.loads(_agent_llm(f"{system}\nGoal: {goal}\n\nSteps so far:\n{steps}\n\nNext step (JSON only):", model))
            bad = 0
        except Exception:
            bad += 1
            obs.append(f"Step {n}: (your reply was not valid JSON - reply with one JSON object)")
            if bad >= 3:
                return "The agent couldn't produce a valid plan. Try a simpler goal or a stronger agent model."
            continue
        if d.get("final"):
            return f"{d['final']}\n\n(Agent finished after {n - 1} step(s).)"
        action, args = d.get("action", ""), d.get("args") or {}
        t = tools.TOOLS.get(action)
        if not t or action in AGENT_BLOCKED or not isinstance(args, dict):
            obs.append(f"Step {n}: '{action}' is not an available tool.")
            continue
        key = (action, json.dumps(args, sort_keys=True, default=str))
        if key in seen:
            obs.append(f"Step {n}: you already did {action} {args}. Do something different or give the final answer.")
            continue
        seen.add(key)
        params = inspect.signature(t["fn"]).parameters
        args = {k: v for k, v in args.items() if k in params}
        print(f"[Agent step {n}] {action} {args}  ({d.get('thought', '')[:80]})")
        logging.info("AGENT %s %s", action, args)
        if t["destructive"] and not tools.ALLOW_DESTRUCTIVE:
            result = "Blocked: destructive actions are disabled."
        elif t["dangerous"] and CONFIRM(f"Amvi (agent): '{action}' {args} is a risky action. "
                                         "Confirm? (yes/no): ").strip().lower() != "yes":
            result = "The user denied this action."
        else:
            try:
                result = str(t["fn"](**args))[:700]
            except TypeError as e:
                result = f"Bad arguments: {e}"
            except Exception as e:
                result = f"Error: {e}"
        obs.append(f"Step {n}: {action} {args} -> {result}")
        time.sleep(0.5)
    return f"I stopped after {MAX_STEPS} steps. Progress so far:\n" + "\n".join(obs)


# =============== 13. CODING MENU (numbered options) ===============
def _cli_menu(title, options):
    print(f"\n{title}")
    for i, o in enumerate(options, 1):
        print(f"  {i}. {o}")
    while True:
        a = input("Choose a number: ").strip()
        if a.isdigit() and 1 <= int(a) <= len(options):
            return int(a)
        print("Please type a number from the list.")


MENU = _cli_menu        # the GUI replaces this with a pop-up of buttons; returns 1-based choice, 0 = closed


def _open_project():
    keys = [k for k in _mem() if "project" in k]
    if not keys:
        return "No saved project yet. Say: remember that my amvi project is at D:\\path"
    if len(keys) == 1:
        key = keys[0]
    else:
        c = MENU("Open which project?", keys + ["Back"])
        if c == 0 or c > len(keys):
            return "Cancelled."
        key = keys[c - 1]
    try:
        return tools.open_saved(key)
    except Exception as e:
        return f"Couldn't open {key}: {e}"


@tool("Coding setup menu: choose VS Code, a project, or a web search")
def coding_menu():
    count = 0
    while True:
        c = MENU("Coding setup", ["VS Code", "Project", "Search from web", "Exit"])
        if c in (0, 4):
            break
        if c == 1:
            msg = tools.open_app("vscode")
        elif c == 2:
            msg = _open_project()
        else:
            q = INPUT("What do you want to search for? ").strip()
            if not q:
                continue
            msg = tools.search_site("google", q)
        count += 1
        print(f"  -> {msg}")
    return f"Coding menu closed. ({count} action{'s' if count != 1 else ''})"


# =============== 14. OPEN-APP MENU AND WEB-SEARCH MENU ===============
def _cli_pick(title, items):
    print(f"\n{title}: " + ", ".join(items))
    while True:
        q = input("Type the name, or part of it (blank to cancel): ").strip().lower()
        if not q:
            return ""
        exact = [i for i in items if i.lower() == q]
        hits = [i for i in items if q in i.lower()]
        if exact:
            return exact[0]
        if len(hits) == 1:
            return hits[0]
        print("Matches: " + ", ".join(hits) if hits else "No match, try again.")


def _cli_file(title):
    return input(f"{title} - full path to the .exe (blank to cancel): ").strip()


PICK = _cli_pick          # the GUI replaces these with a search window and a file picker
PICKFILE = _cli_file


@tool("Open-app menu: pick an available app (with search) or add a new app")
def app_menu():
    count = 0
    while True:
        c = MENU("Open app", ["Available apps", "Add app", "Exit"])
        if c in (0, 3):
            break
        if c == 1:
            name = PICK("Available apps", sorted(APPS))
            if not name:
                continue
            msg = tools.open_app(name)
        else:
            path = PICKFILE("Choose the app to add")
            if not path:
                continue
            default = Path(path.strip('"')).stem.lower()
            name = INPUT(f"Name for this app (blank = {default}): ").strip() or default
            msg = add_app(name, path)
        count += 1
        print(f"  -> {msg}")
    return f"App menu closed. ({count} action{'s' if count != 1 else ''})"


WEB_SITES = [("Google", "google"), ("YouTube (search)", "youtube"), ("GitHub", "github"),
             ("Stack Overflow", "stackoverflow"), ("W3Schools", "w3schools"), ("HackerRank", "hackerrank")]


@tool("Web search menu: choose a site, then type what to search for")
def web_menu():
    labels = [n for n, _ in WEB_SITES] + ["Play on YouTube", "Exit"]
    count = 0
    while True:
        c = MENU("Web search", labels)
        if c in (0, len(labels)):
            break
        q = INPUT("What do you want to search for? ").strip()
        if not q:
            continue
        msg = tools.search_site(WEB_SITES[c - 1][1], q) if c <= len(WEB_SITES) else play_youtube(q)
        count += 1
        print(f"  -> {msg}")
    return f"Web search closed. ({count} search{'es' if count != 1 else ''})"


# =============== 15. WINDOWS, CLIPBOARD, DATA ANALYSIS, WEB PAGES (v11) ===============
@tool("Control the active window: minimize, maximize, restore, switch, snap_left, snap_right, show_desktop",
      '{"action": str}')
def window_action(action):
    if os.name != "nt":
        return "Window control works on Windows only."
    import ctypes
    a = action.lower().strip().replace(" ", "_")
    if a in ("minimize", "minimise", "maximize", "maximise", "restore"):
        hwnd = ctypes.windll.user32.GetForegroundWindow()
        if not hwnd:
            return "No active window."
        ctypes.windll.user32.ShowWindow(hwnd, 6 if a.startswith("minim") else 3 if a.startswith("maxim") else 9)
        return f"Done: {a} window."
    import pyautogui
    keys = {"switch": ("alt", "tab"), "snap_left": ("win", "left"), "snap_right": ("win", "right"),
            "show_desktop": ("win", "d")}
    if a not in keys:
        return "Try: minimize, maximize, restore, switch, snap_left, snap_right, show_desktop."
    pyautogui.hotkey(*keys[a])
    return f"Done: {a}."


@tool("Close the active window", dangerous=True)
def close_window():
    if os.name != "nt":
        return "Window control works on Windows only."
    import ctypes
    hwnd = ctypes.windll.user32.GetForegroundWindow()
    if not hwnd:
        return "No active window."
    ctypes.windll.user32.PostMessageW(hwnd, 0x0010, 0, 0)       # WM_CLOSE: asks the app to close politely
    return "Asked the active window to close."


def _clipboard():
    r = subprocess.run(["powershell", "-NoProfile", "-Command",
                        "[Console]::OutputEncoding=[Text.Encoding]::UTF8; Get-Clipboard -Raw"],
                       capture_output=True, text=True, encoding="utf-8", timeout=10)
    return r.stdout.strip()


@tool("Read the text currently on the clipboard")
def read_clipboard():
    text = _clipboard()
    return text[:3000] if text else "The clipboard is empty (or doesn't hold text)."


@tool("Explain, summarize, fix or translate the text on the clipboard", '{"request": str}')
def ask_clipboard(request="Explain this"):
    text = _clipboard()
    if not text:
        return "The clipboard is empty (or doesn't hold text)."
    prompt = f"{request.strip() or 'Explain this'}\n\nClipboard text:\n{text[:6000]}"
    if re.search(r"summar|translate", request, re.I):
        return _generate(prompt, model_for("study", STUDY_MODEL))
    return _code_llm(prompt)


ANALYSIS_DIR = WORKSPACE / "analysis"
ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)


@tool("Analyze a CSV/Excel dataset: statistics, missing values, duplicates and charts", '{"path": str}')
def analyze_csv(path):
    try:
        import pandas as pd
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as e:
        return f"Missing package ({e.name}). Run: pip install pandas matplotlib openpyxl"
    import warnings
    warnings.filterwarnings("ignore")
    p = _find(path, WORKSPACE, STUDY_DIR)
    if not p.exists():
        return f"File not found: {p}"
    df = pd.read_csv(p) if p.suffix.lower() == ".csv" else pd.read_excel(p)
    out = [f"File: {p.name}", f"Rows: {len(df)}   Columns: {df.shape[1]}",
           "Columns: " + ", ".join(f"{c} ({t})" for c, t in df.dtypes.astype(str).items())]
    miss = df.isna().sum()
    miss = miss[miss > 0]
    out.append("Missing values: " + (", ".join(f"{c}={n}" for c, n in miss.items()) if len(miss) else "none"))
    out.append(f"Duplicate rows: {int(df.duplicated().sum())}")
    num = df.select_dtypes("number")
    if not num.empty:
        out.append("Statistics:\n" + num.describe().T[["mean", "std", "min", "max"]].round(2).to_string())

    charts = []
    if not num.empty:
        cols = list(num.columns[:6])
        fig, axes = plt.subplots(1, len(cols), figsize=(4 * len(cols), 3))
        for ax, c in zip([axes] if len(cols) == 1 else axes, cols):
            ax.hist(num[c].dropna(), bins=20, color="#00a8cc"); ax.set_title(c)
        fig.tight_layout()
        f = ANALYSIS_DIR / f"{p.stem}_histograms.png"; fig.savefig(f); plt.close(fig); charts.append(f)
        date_col = None
        for c in df.columns:
            if not pd.api.types.is_numeric_dtype(df[c]):
                parsed = pd.to_datetime(df[c], errors="coerce")
                if parsed.notna().mean() > 0.8:
                    date_col, df[c] = c, parsed; break
        if date_col is not None:
            ts = df.dropna(subset=[date_col]).sort_values(date_col)
            fig, ax = plt.subplots(figsize=(9, 3.5))
            ax.plot(ts[date_col], ts[num.columns[0]], color="#00a8cc"); ax.set_title(f"{num.columns[0]} over time")
            fig.tight_layout()
            f = ANALYSIS_DIR / f"{p.stem}_trend.png"; fig.savefig(f); plt.close(fig); charts.append(f)
    report = "\n".join(out)
    try:
        report += "\n\nInsights:\n" + _generate("Give 3-4 short plain-language insights about this dataset "
                                                 "summary:\n" + report[:3000], model_for("study", STUDY_MODEL))
    except Exception:
        pass
    saved = ANALYSIS_DIR / f"{p.stem}_report.txt"
    saved.write_text(report, encoding="utf-8")
    if charts and hasattr(os, "startfile"):
        os.startfile(str(charts[0]))
    return report + f"\n\n[Report: {saved}" + "".join(f"\n Chart: {c}" for c in charts) + "]"


@tool("Fetch a web page and summarize it", '{"url": str}')
def summarize_url(url):
    import html as htmlmod
    if not re.match(r"https?://", url):
        return "Give me a full link starting with http:// or https://"
    try:
        r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=15)
        r.raise_for_status()
    except Exception as e:
        return f"Couldn't open that page: {e}"
    raw = r.text
    title = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw)
    body = re.sub(r"(?is)<(script|style|noscript|svg|nav|footer|header)[^>]*>.*?</\1>", " ", raw)
    text = htmlmod.unescape(re.sub(r"\s+", " ", re.sub(r"(?s)<[^>]+>", " ", body))).strip()
    if len(text) < 200:
        return "That page has almost no readable text (it may need JavaScript or a login)."
    head = htmlmod.unescape(title.group(1).strip()) if title else url
    return _generate(f"Summarize this web page in 5-7 short bullet points.\n\nTitle: {head}\n\n{text[:8000]}",
                     model_for("study", STUDY_MODEL))


# =============== 16. PERSONALITY, DAILY BRIEFING, START WITH WINDOWS (v12) ===============
@tool("Choose Amvi's speaking style: friendly, professional, funny, concise or mentor", '{"style": str}')
def set_personality(style=""):
    s = style.strip().lower()
    if s not in tools.PERSONALITIES:
        try:
            cur = json.loads(tools.SETTINGS_FILE.read_text()).get("personality", "friendly")
        except Exception:
            cur = "friendly"
        return f"Styles: {', '.join(tools.PERSONALITIES)}. Now: {cur}. Say, for example: personality funny"
    data = json.loads(tools.SETTINGS_FILE.read_text()) if tools.SETTINGS_FILE.exists() else {}
    data["personality"] = s
    tools.SETTINGS_FILE.write_text(json.dumps(data, indent=2))
    return f"Okay, my style is now {s} ({tools.PERSONALITIES[s]})."


@tool("Daily briefing: date and time, reminders, deadlines, battery")
def daily_briefing():
    import psutil
    now = datetime.now()
    greet = "Good morning" if now.hour < 12 else "Good afternoon" if now.hour < 17 else "Good evening"
    lines = [f"{greet}! It's {now.strftime('%A, %d %B %Y, %I:%M %p')}."]
    rem = sorted((r for r in _load_rem() if not r["done"]), key=lambda r: r["due"])[:3]
    lines.append("Reminders: " + ("; ".join(f"{datetime.fromisoformat(r['due']).strftime('%d %b %I:%M %p')} - {r['text']}"
                                            for r in rem) if rem else "none pending."))
    soon = []
    for d in sorted(tools._deadlines(), key=lambda x: x["date"]):
        days = (datetime.strptime(d["date"], "%Y-%m-%d").date() - now.date()).days
        if days <= 14:
            soon.append(f"{d['task']} ({'overdue' if days < 0 else 'today' if days == 0 else f'in {days} days'})")
    lines.append("Deadlines: " + ("; ".join(soon) if soon else "nothing in the next 2 weeks."))
    b = psutil.sensors_battery()
    if b:
        lines.append(f"Battery: {b.percent:.0f}% ({'charging' if b.power_plugged else 'on battery'}).")
    return "\n".join(lines)


STARTUP_LNK = (Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Windows" / "Start Menu" / "Programs"
               / "Startup" / "Amvi.lnk")


@tool("Start Amvi automatically in the tray when Windows starts: on, off or status", '{"mode": str}')
def startup_launch(mode="status"):
    if os.name != "nt":
        return "This works on Windows only."
    m = mode.strip().lower()
    if m in ("", "status"):
        return "Amvi starts with Windows." if STARTUP_LNK.exists() else "Amvi does not start with Windows."
    if m in ("on", "enable", "yes"):
        vbs = Path(__file__).resolve().parent / "start_amvi_tray.vbs"
        if not vbs.exists():
            return f"I can't find {vbs.name} next to Amvi."
        if CONFIRM("Amvi: Start with Windows (hidden in the tray)? Confirm? (yes/no): ").strip().lower() != "yes":
            return "Cancelled."
        ps = (f"$s=(New-Object -ComObject WScript.Shell).CreateShortcut('{STARTUP_LNK}');"
              f"$s.TargetPath='{vbs}';$s.WorkingDirectory='{vbs.parent}';$s.Save()")
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, timeout=20)
        return ("Done. Amvi will start hidden in the tray when Windows starts." if STARTUP_LNK.exists()
                else "I couldn't create the startup shortcut.")
    if m in ("off", "disable", "no"):
        if CONFIRM("Amvi: Stop starting with Windows? Confirm? (yes/no): ").strip().lower() != "yes":
            return "Cancelled."
        try:
            STARTUP_LNK.unlink()
        except FileNotFoundError:
            pass
        return "Amvi will no longer start with Windows."
    return "Say: start with windows | don't start with windows | startup status"
